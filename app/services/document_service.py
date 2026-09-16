import os
from typing import List, Tuple, Dict, Any

try:
    import pymupdf as fitz
except ImportError:
    try:
        import fitz
    except ImportError:
        fitz = None

try:
    import docx
except ImportError:
    docx = None

class ExtractedDocument:
    def __init__(self, filename: str, file_type: str, file_size: int, pages: List[Tuple[int, str]], total_pages: int):
        self.filename = filename
        self.file_type = file_type
        self.file_size = file_size
        self.pages = pages  # List of (page_number, text)
        self.total_pages = total_pages

class DocumentService:
    MAX_FILE_SIZE_BYTES = 25 * 1024 * 1024  # 25 MB
    ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt"}

    @classmethod
    def validate_file(cls, filename: str, file_size: int) -> Tuple[bool, str]:
        ext = os.path.splitext(filename)[1].lower()
        if ext not in cls.ALLOWED_EXTENSIONS:
            return False, f"Unsupported file format '{ext}'. Supported formats: PDF, DOCX, TXT."
        if file_size > cls.MAX_FILE_SIZE_BYTES:
            return False, f"File size exceeds maximum threshold of 25MB (File size: {file_size / (1024*1024):.2f}MB)."
        return True, ""

    @classmethod
    def extract_text(cls, file_path: str, filename: str, file_size: int) -> ExtractedDocument:
        ext = os.path.splitext(filename)[1].lower()
        
        if ext == ".pdf":
            return cls._extract_pdf(file_path, filename, file_size)
        elif ext == ".docx":
            return cls._extract_docx(file_path, filename, file_size)
        elif ext == ".txt":
            return cls._extract_txt(file_path, filename, file_size)
        else:
            raise ValueError(f"Unsupported file format: {ext}")

    @classmethod
    def _extract_pdf(cls, file_path: str, filename: str, file_size: int) -> ExtractedDocument:
        pages: List[Tuple[int, str]] = []
        if fitz is not None:
            doc = fitz.open(file_path)
            total_pages = len(doc)
            for i, page in enumerate(doc):
                text = page.get_text("text").strip()
                if text:
                    pages.append((i + 1, text))
            doc.close()
        else:
            # Fallback if PyMuPDF fitz not yet loaded
            with open(file_path, "rb") as f:
                raw_bytes = f.read()
            raw_text = "".join([chr(b) for b in raw_bytes if 32 <= b <= 126 or b in (10, 13)])
            pages.append((1, raw_text[:2000]))
            total_pages = 1

        return ExtractedDocument(filename, "pdf", file_size, pages, len(pages) or 1)

    @classmethod
    def _extract_docx(cls, file_path: str, filename: str, file_size: int) -> ExtractedDocument:
        pages: List[Tuple[int, str]] = []
        if docx is not None:
            doc = docx.Document(file_path)
            full_text = []
            for para in doc.paragraphs:
                if para.text.strip():
                    full_text.append(para.text.strip())
            
            current_page_text = []
            word_count = 0
            current_page_num = 1
            
            for para in full_text:
                words = para.split()
                word_count += len(words)
                current_page_text.append(para)
                if word_count >= 400:
                    pages.append((current_page_num, "\n\n".join(current_page_text)))
                    current_page_text = []
                    word_count = 0
                    current_page_num += 1
                    
            if current_page_text:
                pages.append((current_page_num, "\n\n".join(current_page_text)))
        else:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                content = f.read()
            pages.append((1, content))
            
        return ExtractedDocument(filename, "docx", file_size, pages, len(pages) or 1)

    @classmethod
    def _extract_txt(cls, file_path: str, filename: str, file_size: int) -> ExtractedDocument:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            content = f.read()
            
        paragraphs = [p.strip() for p in content.split("\n\n") if p.strip()]
        pages: List[Tuple[int, str]] = []
        current_page_text = []
        word_count = 0
        current_page_num = 1
        
        for para in paragraphs:
            words = para.split()
            word_count += len(words)
            current_page_text.append(para)
            if word_count >= 500:
                pages.append((current_page_num, "\n\n".join(current_page_text)))
                current_page_text = []
                word_count = 0
                current_page_num += 1
                
        if current_page_text:
            pages.append((current_page_num, "\n\n".join(current_page_text)))
        elif not pages and content:
            pages.append((1, content))
            
        return ExtractedDocument(filename, "txt", file_size, pages, len(pages) or 1)
